import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from ilmb_sync.core import ModuleRegistry, WorkbookValidator


def workbook(rows, duplicate_header=False):
    wb = Workbook(); ws = wb.active; ws.title = "02_NUTRIENTS"
    ws.append(["Record_ID", "Name" if not duplicate_header else "Record_ID"])
    for row in rows: ws.append(row)
    out = io.BytesIO(); wb.save(out); return out.getvalue()


class ValidatorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        path = Path(self.temp.name) / "modules.json"
        path.write_text(json.dumps({"modules": [{"pattern": "*Nutrition*.xlsx", "system_id": "SYS-003"}]}))
        self.validator = WorkbookValidator(ModuleRegistry(path))

    def tearDown(self):
        self.temp.cleanup()

    def test_real_rows_stage_and_placeholder_excluded(self):
        result = self.validator.validate_bytes("Nutrition_Food_Medicine.xlsx", workbook([
            ["N-001", "Protein"], ["Example nutrient", "Example"]]))
        self.assertTrue(result.valid)
        self.assertEqual(len(result.rows), 1)
        self.assertEqual(result.rows[0].source_key, "N-001")
        self.assertTrue(any(x.code == "PLACEHOLDER_ROW" for x in result.issues))

    def test_duplicate_keys_are_rejected(self):
        result = self.validator.validate_bytes("Nutrition.xlsx", workbook([
            ["N-001", "Protein"], ["N-001", "Protein"]]))
        self.assertFalse(result.valid)
        self.assertTrue(any(x.code == "DUPLICATE_IDENTICAL_ROW" for x in result.issues))

    def test_repeated_parent_key_gets_composite(self):
        result = self.validator.validate_bytes("Nutrition.xlsx", workbook([
            ["N-001", "Protein"], ["N-001", "Protein revised"]]))
        self.assertTrue(result.valid)
        self.assertEqual(len(result.rows), 2)
        self.assertTrue(any(x.code == "COMPOSITE_SOURCE_KEY" for x in result.issues))

    def test_duplicate_headers_are_rejected(self):
        result = self.validator.validate_bytes("Nutrition.xlsx", workbook([
            ["N-001", "Protein"]], duplicate_header=True))
        self.assertFalse(result.valid)
        self.assertTrue(any(x.code == "DUPLICATE_HEADER" for x in result.issues))

    def test_unregistered_file_is_rejected(self):
        result = self.validator.validate_bytes("Unknown.xlsx", workbook([["X", "Y"]]))
        self.assertFalse(result.valid)
        self.assertEqual(result.issues[0].code, "UNREGISTERED_FILE")

    def test_formula_only_template_row_is_excluded(self):
        wb = Workbook(); ws = wb.active; ws.title = "02_NUTRIENTS"
        ws.append(["Record_ID", "Derived"]); ws.append(["=IF(A1=\"\",\"\",A1)", "=1+1"])
        out = io.BytesIO(); wb.save(out)
        result = self.validator.validate_bytes("Nutrition.xlsx", out.getvalue())
        self.assertFalse(result.valid)
        self.assertEqual(len(result.rows), 0)
        self.assertTrue(any(x.code == "NO_DATA_ROWS" for x in result.issues))

    def test_configured_sheets_and_composite_keys(self):
        path = Path(self.temp.name) / "modules.json"
        path.write_text(json.dumps({"modules": [{
            "pattern": "*Nutrition*.xlsx", "system_id": "SYS-003",
            "include_sheets": ["12_IFCT_VALUES"],
            "sheet_keys": {"12_IFCT_VALUES": ["ingredient_key", "component_key"]}
        }]}))
        validator = WorkbookValidator(ModuleRegistry(path))
        wb = Workbook(); control = wb.active; control.title = "99_DATA_MANIFEST"
        control.append(["Sheet", "Rows"]); control.append(["12_IFCT_VALUES", 2])
        ws = wb.create_sheet("12_IFCT_VALUES")
        ws.append(["ingredient_key", "component_key", "amount"])
        ws.append(["ifct_a001", "protein", 12.3]); ws.append(["ifct_a001", "fat", 7.1])
        out = io.BytesIO(); wb.save(out)
        result = validator.validate_bytes("Nutrition.xlsx", out.getvalue())
        self.assertTrue(result.valid)
        self.assertEqual([r.source_key for r in result.rows], ["ifct_a001::protein", "ifct_a001::fat"])
        self.assertNotIn("99_DATA_MANIFEST", result.sheet_counts)


if __name__ == "__main__":
    unittest.main()

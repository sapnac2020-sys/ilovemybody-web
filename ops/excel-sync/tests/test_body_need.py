from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from ilmb_sync.body_need import (
    FormulaError,
    FormulaInputSpec,
    assess_formula_execution,
    duplicate_candidates,
    evaluate_formula,
    export_formula_master_xlsx,
    formula_symbols,
    normalized_expression,
)


class BodyNeedTests(unittest.TestCase):
    def test_safe_arithmetic(self):
        self.assertEqual(str(evaluate_formula("target-measured", {"target": 10, "measured": 4})), "6")
        self.assertEqual(formula_symbols("a + b*2"), {"a", "b"})
        self.assertEqual(normalized_expression("a+b"), normalized_expression("a + b"))

    def test_disallows_calls(self):
        with self.assertRaises(FormulaError):
            evaluate_formula("__import__('os').system('id')", {})

    def test_unverified_target_blocks(self):
        result = assess_formula_execution(
            expression="target-measured",
            formula_status="APPROVED",
            unit_checked=True,
            input_specs=[
                FormulaInputSpec("target", "TARGET", True, "mg"),
                FormulaInputSpec("measured", "MEASURED", True, "mg"),
            ],
            observations={
                "target": {"value": 10, "unit": "mg", "verified_source": False},
                "measured": {"value": 4, "unit": "mg", "verified_source": True},
            },
        )
        self.assertEqual(result["status"], "BLOCKED_UNVERIFIED_TARGET")

    def test_verified_target_computes(self):
        result = assess_formula_execution(
            expression="target-measured",
            formula_status="VERIFIED",
            unit_checked=True,
            input_specs=[FormulaInputSpec("target", "TARGET", True, "mg"), FormulaInputSpec("measured", "MEASURED", True, "mg")],
            observations={
                "target": {"value": 10, "unit": "mg", "verified_source": True, "source": "governed"},
                "measured": {"value": 4, "unit": "mg", "verified_source": True, "source": "lab"},
            },
        )
        self.assertEqual(result["status"], "COMPUTED")
        self.assertEqual(result["value"], "6")

    def test_duplicate_signals(self):
        params = [
            {"parameter_id": 1, "canonical_name": "Serum sodium", "canonical_ucum_unit": "mmol/L", "source_system": "A", "source_record_key": "X"},
            {"parameter_id": 2, "canonical_name": "Serum Sodium", "canonical_ucum_unit": "mmol/L", "source_system": "B", "source_record_key": "Y"},
        ]
        identifiers = [
            {"parameter_id": 1, "identifier_system": "LOINC", "identifier_code": "2951-2", "verification_status": "APPROVED"},
            {"parameter_id": 2, "identifier_system": "LOINC", "identifier_code": "2951-2", "verification_status": "APPROVED"},
        ]
        reasons = {x["reason"] for x in duplicate_candidates(params, identifiers)}
        self.assertIn("SAME_IDENTIFIER", reasons)
        self.assertIn("NORMALIZED_NAME_UNIT", reasons)

    def test_export_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "formula-master.xlsx"
            export_formula_master_xlsx(
                path,
                parameters=[{"parameter_id": 1, "parameter_key": "sodium"}],
                identifiers=[], formulas=[], formula_inputs=[], duplicates=[])
            wb = load_workbook(path, read_only=True)
            self.assertEqual(
                wb.sheetnames,
                ["Parameters", "Identifiers", "Formula_Master", "Formula_Inputs", "Duplicate_Candidates", "Body_Need_Template"],
            )


if __name__ == "__main__":
    unittest.main()

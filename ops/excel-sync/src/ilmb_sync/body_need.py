from __future__ import annotations

import ast
import base64
import hashlib
import json
import re
import urllib.parse
import urllib.request
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook


_ALLOWED_BINOPS = (ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow)
_ALLOWED_UNARYOPS = (ast.UAdd, ast.USub)
_NAME_RE = re.compile(r"[^a-z0-9]+")


class FormulaError(ValueError):
    pass


@dataclass(frozen=True)
class FormulaInputSpec:
    symbol: str
    role: str
    required: bool = True
    expected_unit: str | None = None


def normalize_name(value: str) -> str:
    return _NAME_RE.sub(" ", (value or "").lower()).strip()


def normalized_expression(expression: str) -> str:
    tree = ast.parse(expression, mode="eval")
    _validate_ast(tree)
    return ast.dump(tree, annotate_fields=False, include_attributes=False)


def _validate_ast(node: ast.AST) -> None:
    for item in ast.walk(node):
        if isinstance(item, (ast.Expression, ast.Load, ast.Constant, ast.Name)):
            continue
        if isinstance(item, ast.BinOp) and isinstance(item.op, _ALLOWED_BINOPS):
            continue
        if isinstance(item, ast.UnaryOp) and isinstance(item.op, _ALLOWED_UNARYOPS):
            continue
        if isinstance(item, _ALLOWED_BINOPS + _ALLOWED_UNARYOPS):
            continue
        raise FormulaError(f"Unsupported formula syntax: {type(item).__name__}")


def formula_symbols(expression: str) -> set[str]:
    tree = ast.parse(expression, mode="eval")
    _validate_ast(tree)
    return {item.id for item in ast.walk(tree) if isinstance(item, ast.Name)}


def evaluate_formula(expression: str, values: dict[str, int | float | Decimal]) -> Decimal:
    tree = ast.parse(expression, mode="eval")
    _validate_ast(tree)
    missing = formula_symbols(expression) - set(values)
    if missing:
        raise FormulaError("Missing formula inputs: " + ", ".join(sorted(missing)))

    def walk(node: ast.AST) -> Decimal:
        if isinstance(node, ast.Expression):
            return walk(node.body)
        if isinstance(node, ast.Constant):
            if not isinstance(node.value, (int, float)):
                raise FormulaError("Only numeric constants are allowed")
            return Decimal(str(node.value))
        if isinstance(node, ast.Name):
            return Decimal(str(values[node.id]))
        if isinstance(node, ast.UnaryOp):
            value = walk(node.operand)
            return value if isinstance(node.op, ast.UAdd) else -value
        if isinstance(node, ast.BinOp):
            left, right = walk(node.left), walk(node.right)
            if isinstance(node.op, ast.Add):
                return left + right
            if isinstance(node.op, ast.Sub):
                return left - right
            if isinstance(node.op, ast.Mult):
                return left * right
            if isinstance(node.op, ast.Div):
                return left / right
            if isinstance(node.op, ast.Pow):
                if right != int(right):
                    raise FormulaError("Only integer exponents are supported")
                return left ** int(right)
        raise FormulaError(f"Unsupported formula node: {type(node).__name__}")

    return walk(tree)


def assess_formula_execution(
    *,
    expression: str,
    formula_status: str,
    unit_checked: bool,
    input_specs: Iterable[FormulaInputSpec],
    observations: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    if formula_status not in {"VERIFIED", "APPROVED"}:
        return {"status": "BLOCKED_UNVERIFIED_FORMULA", "reason": "Formula is not verified/approved"}
    if not unit_checked:
        return {"status": "BLOCKED_UNIT_MISMATCH", "reason": "Formula has not passed unit checking"}

    specs = {x.symbol: x for x in input_specs}
    expected = formula_symbols(expression)
    undeclared = expected - set(specs)
    if undeclared:
        return {"status": "BLOCKED_MISSING_INPUT", "reason": "Undeclared symbols: " + ", ".join(sorted(undeclared))}

    values: dict[str, Any] = {}
    provenance: dict[str, Any] = {}
    for symbol in sorted(expected):
        spec = specs[symbol]
        observed = observations.get(symbol)
        if not observed:
            if spec.required:
                return {"status": "BLOCKED_MISSING_INPUT", "reason": f"Missing required input: {symbol}"}
            continue
        if spec.expected_unit and observed.get("unit") != spec.expected_unit:
            return {"status": "BLOCKED_UNIT_MISMATCH", "reason": f"Unit mismatch for {symbol}"}
        if spec.role == "TARGET" and not observed.get("verified_source"):
            return {"status": "BLOCKED_UNVERIFIED_TARGET", "reason": f"Target lacks governed provenance: {symbol}"}
        values[symbol] = observed["value"]
        provenance[symbol] = {k: observed.get(k) for k in ("unit", "source", "identifier_system", "identifier_code", "verified_source")}

    result = evaluate_formula(expression, values)
    return {"status": "COMPUTED", "value": str(result), "inputs": values, "provenance": provenance}


def duplicate_candidates(parameters: Iterable[dict[str, Any]], identifiers: Iterable[dict[str, Any]], formulas: Iterable[dict[str, Any]] = ()) -> list[dict[str, Any]]:
    params = {int(p["parameter_id"]): p for p in parameters}
    found: dict[tuple[int, int, str], dict[str, Any]] = {}

    by_identifier: dict[tuple[str, str], list[int]] = {}
    for row in identifiers:
        if row.get("verification_status") != "APPROVED":
            continue
        by_identifier.setdefault((str(row["identifier_system"]), str(row["identifier_code"])), []).append(int(row["parameter_id"]))
    for key, ids in by_identifier.items():
        unique_ids = sorted(set(ids))
        for i, left in enumerate(unique_ids):
            for right in unique_ids[i + 1:]:
                found[(left, right, "SAME_IDENTIFIER")] = {
                    "left_parameter_id": left,
                    "right_parameter_id": right,
                    "reason": "SAME_IDENTIFIER",
                    "confidence": 1.0,
                    "evidence": {"system": key[0], "code": key[1]},
                }

    by_name_unit: dict[tuple[str, str], list[int]] = {}
    by_source: dict[tuple[str, str], list[int]] = {}
    for pid, p in params.items():
        name_key = (normalize_name(str(p.get("canonical_name", ""))), str(p.get("canonical_ucum_unit") or ""))
        if name_key[0]:
            by_name_unit.setdefault(name_key, []).append(pid)
        if p.get("source_system") and p.get("source_record_key"):
            by_source.setdefault((str(p["source_system"]), str(p["source_record_key"])), []).append(pid)
    for key, ids in by_name_unit.items():
        unique_ids = sorted(set(ids))
        for i, left in enumerate(unique_ids):
            for right in unique_ids[i + 1:]:
                found[(left, right, "NORMALIZED_NAME_UNIT")] = {
                    "left_parameter_id": left,
                    "right_parameter_id": right,
                    "reason": "NORMALIZED_NAME_UNIT",
                    "confidence": 0.85,
                    "evidence": {"normalized_name": key[0], "unit": key[1]},
                }
    for key, ids in by_source.items():
        unique_ids = sorted(set(ids))
        for i, left in enumerate(unique_ids):
            for right in unique_ids[i + 1:]:
                found[(left, right, "SAME_SOURCE_KEY")] = {
                    "left_parameter_id": left,
                    "right_parameter_id": right,
                    "reason": "SAME_SOURCE_KEY",
                    "confidence": 0.98,
                    "evidence": {"source_system": key[0], "source_key": key[1]},
                }

    return sorted(found.values(), key=lambda x: (x["left_parameter_id"], x["right_parameter_id"], x["reason"]))


def formula_duplicate_candidates(formulas: Iterable[dict[str, Any]], formula_inputs: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    formulas_by_id = {int(row["formula_id"]): row for row in formulas}
    inputs_by_formula: dict[int, list[tuple[int, str, str]]] = {}
    for row in formula_inputs:
        fid = int(row["formula_id"])
        inputs_by_formula.setdefault(fid, []).append(
            (int(row["parameter_id"]), str(row.get("role") or ""), str(row.get("expected_ucum_unit") or ""))
        )

    found: dict[tuple[int, int, str], dict[str, Any]] = {}
    by_expression: dict[str, list[int]] = {}
    for fid, row in formulas_by_id.items():
        expression = str(row.get("expression_text") or "").strip()
        if not expression:
            continue
        try:
            normalized = normalized_expression(expression)
        except (FormulaError, SyntaxError, ValueError):
            continue
        by_expression.setdefault(normalized, []).append(fid)

    for normalized, ids in by_expression.items():
        unique_ids = sorted(set(ids))
        digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
        for i, left in enumerate(unique_ids):
            for right in unique_ids[i + 1:]:
                found[(left, right, "NORMALIZED_EXPRESSION")] = {
                    "left_formula_id": left,
                    "right_formula_id": right,
                    "reason": "NORMALIZED_EXPRESSION",
                    "confidence": 0.97,
                    "evidence": {"normalized_expression_sha256": digest},
                }

    by_signature: dict[tuple[Any, tuple[tuple[int, str, str], ...]], list[int]] = {}
    for fid, row in formulas_by_id.items():
        signature = (
            row.get("output_parameter_id"),
            tuple(sorted(inputs_by_formula.get(fid, []))),
        )
        if signature[0] is None and not signature[1]:
            continue
        by_signature.setdefault(signature, []).append(fid)

    for signature, ids in by_signature.items():
        unique_ids = sorted(set(ids))
        for i, left in enumerate(unique_ids):
            for right in unique_ids[i + 1:]:
                found[(left, right, "SAME_OUTPUT_AND_INPUTS")] = {
                    "left_formula_id": left,
                    "right_formula_id": right,
                    "reason": "SAME_OUTPUT_AND_INPUTS",
                    "confidence": 0.90,
                    "evidence": {
                        "output_parameter_id": signature[0],
                        "input_signature": [list(x) for x in signature[1]],
                    },
                }

    return sorted(found.values(), key=lambda x: (x["left_formula_id"], x["right_formula_id"], x["reason"]))


def _json_request(url: str, *, username: str | None = None, password: str | None = None, timeout: int = 30) -> dict[str, Any]:
    req = urllib.request.Request(url, headers={"Accept": "application/json", "User-Agent": "ILoveMyBody/1.0"})
    if username is not None and password is not None:
        token = base64.b64encode(f"{username}:{password}".encode()).decode()
        req.add_header("Authorization", f"Basic {token}")
    with urllib.request.urlopen(req, timeout=timeout) as response:
        return json.loads(response.read().decode("utf-8"))


def loinc_lookup(code: str, *, username: str, password: str, base_url: str = "https://fhir.loinc.org", timeout: int = 30) -> dict[str, Any]:
    query = urllib.parse.urlencode({"system": "http://loinc.org", "code": code})
    return _json_request(f"{base_url.rstrip('/')}/CodeSystem/$lookup?{query}", username=username, password=password, timeout=timeout)


def export_formula_master_xlsx(
    output_path: str | Path,
    *,
    parameters: Iterable[dict[str, Any]],
    identifiers: Iterable[dict[str, Any]],
    formulas: Iterable[dict[str, Any]],
    formula_inputs: Iterable[dict[str, Any]],
    duplicates: Iterable[dict[str, Any]],
    formula_duplicates: Iterable[dict[str, Any]] = (),
) -> Path:
    path = Path(output_path)
    wb = Workbook()
    wb.remove(wb.active)

    datasets = {
        "Parameters": list(parameters),
        "Identifiers": list(identifiers),
        "Formula_Master": list(formulas),
        "Formula_Inputs": list(formula_inputs),
        "Duplicate_Candidates": list(duplicates),
        "Formula_Duplicates": list(formula_duplicates),
    }
    for title, rows in datasets.items():
        ws = wb.create_sheet(title)
        headers: list[str] = []
        for row in rows:
            for key in row:
                if key not in headers:
                    headers.append(key)
        if not headers:
            headers = ["status"]
            rows = [{"status": "NO_ROWS"}]
        ws.append(headers)
        for row in rows:
            ws.append([
                json.dumps(row.get(h), ensure_ascii=False) if isinstance(row.get(h), (dict, list)) else row.get(h)
                for h in headers
            ])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    template = wb.create_sheet("Body_Need_Template")
    template.append(["subject_key", "formula_key", "symbol", "value", "ucum_unit", "identifier_system", "identifier_code", "source", "verified_source"])
    template.freeze_panes = "A2"
    wb.save(path)
    return path

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import Database

EXCEL_MAX_ROWS = 1_048_576
HEADER_FILL = PatternFill("solid", fgColor="E5E7EB")
TITLE_FILL = PatternFill("solid", fgColor="1F2937")


def database_from_env() -> Database:
    return Database(
        host=os.environ["ILMB_DB_HOST"],
        port=int(os.environ.get("ILMB_DB_PORT", "3306")),
        name=os.environ["ILMB_DB_NAME"],
        user=os.environ["ILMB_DB_USER"],
        password=os.environ["ILMB_DB_PASSWORD"],
    )


def safe_sheet_name(raw: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", raw).strip() or "Sheet"
    base = base[:31]
    name = base
    n = 2
    while name.lower() in used:
        suffix = f"_{n}"
        name = (base[: 31 - len(suffix)] + suffix)
        n += 1
    used.add(name.lower())
    return name


def style_header(ws, headers: list[str]) -> None:
    for idx, value in enumerate(headers, start=1):
        c = ws.cell(row=1, column=idx, value=value)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max(1, len(headers)))}1"


def normalize_cell(value):
    if value is None or isinstance(value, (str, int, float, bool, datetime)):
        return value
    if isinstance(value, (bytes, bytearray)):
        return value.hex()
    return str(value)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def save_workbook(wb: Workbook, path: Path) -> dict:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"file": path.name, "bytes": path.stat().st_size, "sha256": sha256_file(path)}


def export_query_sheet(conn, wb: Workbook, name: str, sql: str, params: tuple = ()) -> dict:
    used = {ws.title.lower() for ws in wb.worksheets}
    ws = wb.create_sheet(safe_sheet_name(name, used))
    with conn.cursor() as cur:
        cur.execute(sql, params)
        headers = [d[0] for d in cur.description] if cur.description else []
        style_header(ws, headers)
        count = 0
        while True:
            rows = cur.fetchmany(2000)
            if not rows:
                break
            for row in rows:
                if isinstance(row, dict):
                    ws.append([normalize_cell(row.get(h)) for h in headers])
                else:
                    ws.append([normalize_cell(v) for v in row])
                count += 1
    return {"sheet": ws.title, "rows": count, "columns": len(headers)}


def export_body_need_master(output_root: Path) -> dict:
    db = database_from_env()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    release = output_root / "formula" / "releases" / stamp
    release.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=False)
    wb.remove(wb.active)
    queries = [
        ("Parameters", "SELECT * FROM ilmb_parameter_master ORDER BY parameter_domain,canonical_name"),
        ("Identifiers", "SELECT i.*,p.parameter_key,p.canonical_name FROM ilmb_parameter_identifier i JOIN ilmb_parameter_master p ON p.parameter_id=i.parameter_id ORDER BY p.parameter_key,i.identifier_system,i.identifier_code"),
        ("Formula_Master", "SELECT f.*,p.parameter_key output_parameter_key,p.canonical_name output_parameter_name FROM ilmb_formula_master f LEFT JOIN ilmb_parameter_master p ON p.parameter_id=f.output_parameter_id ORDER BY f.formula_domain,f.formula_name"),
        ("Formula_Inputs", "SELECT i.*,f.formula_key,p.parameter_key,p.canonical_name FROM ilmb_formula_input i JOIN ilmb_formula_master f ON f.formula_id=i.formula_id JOIN ilmb_parameter_master p ON p.parameter_id=i.parameter_id ORDER BY f.formula_key,i.ordinal,i.symbol_name"),
        ("Parameter_Duplicates", "SELECT d.*,l.parameter_key left_key,r.parameter_key right_key FROM ilmb_parameter_duplicate_candidate d JOIN ilmb_parameter_master l ON l.parameter_id=d.left_parameter_id JOIN ilmb_parameter_master r ON r.parameter_id=d.right_parameter_id ORDER BY d.created_at DESC"),
        ("Formula_Duplicates", "SELECT d.*,l.formula_key left_key,r.formula_key right_key FROM ilmb_formula_duplicate_candidate d JOIN ilmb_formula_master l ON l.formula_id=d.left_formula_id JOIN ilmb_formula_master r ON r.formula_id=d.right_formula_id ORDER BY d.created_at DESC"),
        ("Body_Need_Runs", "SELECT * FROM ilmb_body_need_run ORDER BY created_at DESC LIMIT 10000"),
    ]
    sheets = []
    with db.connect() as conn:
        for name, sql in queries:
            sheets.append(export_query_sheet(conn, wb, name, sql))
    file_info = save_workbook(wb, release / "ILMB_Formula_Master_Live.xlsx")
    manifest = {"kind":"body_need_master","generated_at_utc":stamp,"release":str(release),"sheets":sheets,"files":[file_info]}
    (release / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    current = output_root / "formula" / "current"
    if current.is_symlink() or current.exists():
        if current.is_symlink() or current.is_file(): current.unlink()
        else:
            import shutil; shutil.rmtree(current)
    current.symlink_to(release, target_is_directory=True)
    return manifest


def table_inventory(conn) -> tuple[list[dict], list[dict]]:
    with conn.cursor() as cur:
        cur.execute("SELECT TABLE_NAME,TABLE_TYPE,ENGINE,TABLE_ROWS,DATA_LENGTH,INDEX_LENGTH,CREATE_TIME,UPDATE_TIME FROM information_schema.TABLES WHERE TABLE_SCHEMA=DATABASE() ORDER BY TABLE_TYPE,TABLE_NAME")
        rows = cur.fetchall()
        base = [r for r in rows if r["TABLE_TYPE"] == "BASE TABLE"]
        views = [r for r in rows if r["TABLE_TYPE"] == "VIEW"]
        return base, views


def export_table_to_book(conn, wb: Workbook, table: str, used: set[str]) -> dict:
    safe_table = table.replace("`", "``")
    with conn.cursor() as cur:
        cur.execute(f"SELECT * FROM `{safe_table}`")
        headers = [d[0] for d in cur.description] if cur.description else []
        part = 1
        ws = wb.create_sheet(safe_sheet_name(table, used))
        style_header(ws, headers)
        rows_in_sheet = 0
        total = 0
        parts = [ws.title]
        while True:
            batch = cur.fetchmany(1000)
            if not batch: break
            for row in batch:
                if rows_in_sheet >= EXCEL_MAX_ROWS - 1:
                    part += 1
                    ws = wb.create_sheet(safe_sheet_name(f"{table}_{part}", used))
                    style_header(ws, headers)
                    rows_in_sheet = 0
                    parts.append(ws.title)
                ws.append([normalize_cell(row.get(h)) for h in headers])
                rows_in_sheet += 1
                total += 1
    return {"table":table,"rows":total,"columns":len(headers),"sheets":parts}


def export_full_db(output_root: Path, tables_per_workbook: int = 20) -> dict:
    db = database_from_env()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    release = output_root / "db-mirror" / "releases" / stamp
    release.mkdir(parents=True, exist_ok=True)
    workbook_files = []
    table_results = []
    with db.connect() as conn:
        base_tables, views = table_inventory(conn)
        # Control room first.
        control = Workbook(write_only=False)
        ws = control.active; ws.title = "CONTROL_ROOM"
        headers = ["table_name","table_type","engine","estimated_rows","data_length","index_length","create_time","update_time"]
        style_header(ws, headers)
        for r in base_tables + views:
            ws.append([normalize_cell(r.get(k.upper()) if k.upper() in r else r.get(k)) for k in headers])
        vws = control.create_sheet("VIEWS")
        style_header(vws,["view_name","definition"])
        with conn.cursor() as cur:
            for r in views:
                name=r["TABLE_NAME"]
                cur.execute("SELECT VIEW_DEFINITION FROM information_schema.VIEWS WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME=%s",(name,))
                d=cur.fetchone()
                vws.append([name,(d or {}).get("VIEW_DEFINITION")])
        control_info = save_workbook(control, release / "00_CONTROL_ROOM.xlsx")
        workbook_files.append(control_info)

        for start in range(0, len(base_tables), tables_per_workbook):
            batch = base_tables[start:start+tables_per_workbook]
            wb = Workbook(write_only=False)
            wb.remove(wb.active)
            used=set()
            batch_results=[]
            for meta in batch:
                result=export_table_to_book(conn,wb,meta["TABLE_NAME"],used)
                batch_results.append(result); table_results.append(result)
            idx = start // tables_per_workbook + 1
            info=save_workbook(wb, release / f"DB_Mirror_{idx:03d}.xlsx")
            info["tables"]=[x["table"] for x in batch_results]
            workbook_files.append(info)
    manifest={"kind":"full_db_excel_mirror","generated_at_utc":stamp,"base_tables":len(base_tables),"views":len(views),"tables":table_results,"files":workbook_files}
    (release / "manifest.json").write_text(json.dumps(manifest,indent=2,default=str),encoding="utf-8")
    current=output_root / "db-mirror" / "current"
    if current.is_symlink() or current.exists():
        if current.is_symlink() or current.is_file(): current.unlink()
        else:
            import shutil; shutil.rmtree(current)
    current.symlink_to(release,target_is_directory=True)
    return manifest


def write_root_manifest(root: Path, last: dict) -> None:
    payload={"generated_at_utc":datetime.now(timezone.utc).isoformat(),"last_job":last,"private_root":str(root),"public_html_excluded":True}
    root.mkdir(parents=True,exist_ok=True)
    (root / "manifest.json").write_text(json.dumps(payload,indent=2,default=str),encoding="utf-8")


def main(argv: Iterable[str] | None = None) -> int:
    p=argparse.ArgumentParser()
    p.add_argument("mode",choices=["body-master","full-db"])
    p.add_argument("--output-root",default=os.environ.get("ILMB_PRIVATE_EXCEL_ROOT",str(Path.home()/"ilmb-data-exchange")))
    p.add_argument("--tables-per-workbook",type=int,default=20)
    args=p.parse_args(list(argv) if argv is not None else None)
    root=Path(args.output_root).expanduser().resolve()
    result=export_body_need_master(root) if args.mode=="body-master" else export_full_db(root,args.tables_per_workbook)
    write_root_manifest(root,result)
    print(json.dumps(result,default=str))
    return 0

if __name__ == "__main__":
    raise SystemExit(main())

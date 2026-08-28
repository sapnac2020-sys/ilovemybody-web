from __future__ import annotations

import fnmatch
import hashlib
import io
import json
import os
import re
import uuid
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


PLACEHOLDER_RE = re.compile(r"^(example|sample|placeholder|test)(\b|[_ -])", re.I)
CONTROL_SHEETS = re.compile(r"(^00_|readme|dashboard|codebook|audit|setup|control|log$|(?:^|_)qc$)", re.I)


def utcnow() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def stable_hash(*parts: str) -> str:
    return hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()


@dataclass(frozen=True)
class ValidationIssue:
    severity: str
    code: str
    message: str
    sheet_name: str | None = None
    source_row: int | None = None


@dataclass(frozen=True)
class StageRow:
    row_id: str
    system_id: str
    sheet_name: str
    source_row: int
    source_key: str
    row_hash: str
    payload: dict[str, Any]


@dataclass
class ValidationResult:
    file_name: str
    system_id: str
    content_hash: str
    rows: list[StageRow]
    issues: list[ValidationIssue]
    sheet_counts: dict[str, int]

    @property
    def valid(self) -> bool:
        return not any(i.severity == "ERROR" for i in self.issues)

    def summary(self) -> dict[str, Any]:
        return {
            "file_name": self.file_name,
            "system_id": self.system_id,
            "content_hash": self.content_hash,
            "valid": self.valid,
            "row_count": len(self.rows),
            "error_count": sum(i.severity == "ERROR" for i in self.issues),
            "warning_count": sum(i.severity == "WARNING" for i in self.issues),
            "sheet_counts": self.sheet_counts,
            "issues": [asdict(x) for x in self.issues],
        }


class ModuleRegistry:
    def __init__(self, config_path: str | Path):
        self.path = Path(config_path)
        self.data = json.loads(self.path.read_text(encoding="utf-8"))

    def match(self, file_name: str) -> dict[str, Any] | None:
        matches = [m for m in self.data["modules"] if fnmatch.fnmatch(file_name.lower(), m["pattern"].lower())]
        if len(matches) > 1:
            raise ValueError(f"Ambiguous module rules for {file_name}: {[m['pattern'] for m in matches]}")
        return matches[0] if matches else None


class WorkbookValidator:
    def __init__(self, registry: ModuleRegistry):
        self.registry = registry

    def validate_bytes(self, file_name: str, content: bytes) -> ValidationResult:
        rule = self.registry.match(file_name)
        content_hash = sha256_bytes(content)
        if not rule:
            return ValidationResult(file_name, "UNREGISTERED", content_hash, [], [
                ValidationIssue("ERROR", "UNREGISTERED_FILE", "No module registry rule matches this workbook")
            ], {})
        issues: list[ValidationIssue] = []
        rows: list[StageRow] = []
        counts: dict[str, int] = {}
        try:
            wb = load_workbook(io.BytesIO(content), data_only=False, read_only=True, keep_links=False)
        except Exception as exc:
            return ValidationResult(file_name, rule["system_id"], content_hash, [], [
                ValidationIssue("ERROR", "INVALID_XLSX", f"Workbook cannot be opened: {type(exc).__name__}: {exc}")
            ], {})

        seen_sources: dict[tuple[str, str], str] = {}
        include_sheets = set(rule.get("include_sheets", []))
        sheet_keys = rule.get("sheet_keys", {})
        for ws in wb.worksheets:
            if CONTROL_SHEETS.search(ws.title):
                continue
            if include_sheets and ws.title not in include_sheets:
                continue
            values = ws.iter_rows(values_only=True)
            header_row = None
            header_num = 0
            for idx, candidate in enumerate(values, start=1):
                nonblank = [v for v in candidate if v is not None and str(v).strip()]
                if len(nonblank) >= 2:
                    header_row, header_num = candidate, idx
                    break
            if header_row is None:
                continue
            headers = self._headers(header_row, ws.title, issues, header_num)
            if not headers:
                continue
            sheet_count = 0
            row_gate = rule.get("row_gates", {}).get(ws.title, {})
            for source_row, vals in enumerate(values, start=header_num + 1):
                payload = {headers[i]: self._cell(vals[i]) for i in range(min(len(headers), len(vals)))}
                payload = {k: v for k, v in payload.items() if v not in (None, "")}
                for field in rule.get("numeric_fields", {}).get(ws.title, []):
                    if field in payload:
                        payload[field] = self._number(payload[field])
                if not payload:
                    continue
                # Preformatted enterprise templates often copy formulas hundreds
                # of rows below the real data. Formula-only rows are scaffolding,
                # not records, and must never inflate database counts.
                if not any(not isinstance(value, dict) for value in payload.values()):
                    continue
                missing_required = [field for field in row_gate.get("required", []) if payload.get(field) in (None, "")]
                if missing_required:
                    issues.append(ValidationIssue("WARNING", "ROW_EXCLUDED_MISSING_REQUIRED", f"Row excluded; missing required fields: {', '.join(missing_required)}", ws.title, source_row))
                    continue
                invalid_numeric = [field for field in row_gate.get("numeric", []) if not isinstance(payload.get(field), (int, float))]
                if invalid_numeric:
                    issues.append(ValidationIssue("WARNING", "ROW_EXCLUDED_NONNUMERIC", f"Row excluded; numeric values required for: {', '.join(invalid_numeric)}", ws.title, source_row))
                    continue
                invalid_negative = [field for field in row_gate.get("nonnegative", []) if isinstance(payload.get(field), (int, float)) and payload[field] < 0]
                if invalid_negative:
                    issues.append(ValidationIssue("WARNING", "ROW_EXCLUDED_NEGATIVE_VALUE", f"Row excluded; negative values are invalid for: {', '.join(invalid_negative)}", ws.title, source_row))
                    continue
                source_key = self._source_key(payload, source_row, sheet_keys.get(ws.title))
                if PLACEHOLDER_RE.search(source_key):
                    issues.append(ValidationIssue("WARNING", "PLACEHOLDER_ROW", "Placeholder/example row excluded", ws.title, source_row))
                    continue
                row_json = canonical_json(payload)
                row_hash = stable_hash(rule["system_id"], ws.title, source_key, row_json)
                identity = (ws.title, source_key)
                if identity in seen_sources:
                    if seen_sources[identity] == row_hash:
                        issues.append(ValidationIssue("ERROR", "DUPLICATE_IDENTICAL_ROW", f"Exact duplicate row for source key: {source_key}", ws.title, source_row))
                        continue
                    # Link/observation tables legitimately repeat a parent code.
                    # Preserve every distinct row with an explicit composite key and
                    # flag it for later replacement by a governed workbook row ID.
                    source_key = f"{source_key}::{stable_hash(row_json)[:16]}"
                    issues.append(ValidationIssue("WARNING", "COMPOSITE_SOURCE_KEY", "Repeated natural key; deterministic composite key assigned", ws.title, source_row))
                    row_hash = stable_hash(rule["system_id"], ws.title, source_key, row_json)
                    identity = (ws.title, source_key)
                    if identity in seen_sources:
                        issues.append(ValidationIssue("ERROR", "DUPLICATE_COMPOSITE_KEY", f"Composite key collision: {source_key}", ws.title, source_row))
                        continue
                seen_sources[identity] = row_hash
                rows.append(StageRow(
                    row_id=stable_hash(content_hash, ws.title, str(source_row), source_key),
                    system_id=rule["system_id"], sheet_name=ws.title, source_row=source_row,
                    source_key=source_key, row_hash=row_hash, payload=payload,
                ))
                sheet_count += 1
            if sheet_count:
                counts[ws.title] = sheet_count
        if not rows:
            issues.append(ValidationIssue("ERROR", "NO_DATA_ROWS", "No importable non-placeholder rows were found"))
        wb.close()
        return ValidationResult(file_name, rule["system_id"], content_hash, rows, issues, counts)

    @staticmethod
    def _headers(cells: Iterable[Any], sheet: str, issues: list[ValidationIssue], row: int) -> list[str]:
        result: list[str] = []
        seen: set[str] = set()
        for idx, value in enumerate(cells, start=1):
            name = str(value).strip().lstrip("\ufeff") if value is not None else f"_column_{idx}"
            if not name:
                name = f"_column_{idx}"
            if name in seen:
                issues.append(ValidationIssue("ERROR", "DUPLICATE_HEADER", f"Duplicate header: {name}", sheet, row))
                return []
            seen.add(name)
            result.append(name)
        return result

    @staticmethod
    def _cell(value: Any) -> Any:
        if isinstance(value, datetime):
            return value.isoformat()
        if isinstance(value, str):
            value = value.strip()
            if value.startswith("="):
                return {"formula": value}
        return value

    @staticmethod
    def _number(value: Any) -> Any:
        if not isinstance(value, str):
            return value
        text = value.strip()
        if not re.fullmatch(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][-+]?\d+)?", text):
            return value
        number = float(text)
        return int(number) if number.is_integer() and "e" not in text.lower() and "." not in text else number

    @staticmethod
    def _source_key(payload: dict[str, Any], source_row: int, configured_keys: list[str] | None = None) -> str:
        if configured_keys:
            missing = [key for key in configured_keys if key not in payload or isinstance(payload[key], dict) or not str(payload[key]).strip()]
            if not missing:
                return "::".join(str(payload[key]).strip() for key in configured_keys)
        # Prefer an actual row identifier. Evidence/source IDs are foreign keys and
        # must not be mistaken for the row's own identity.
        for key, value in payload.items():
            normalized = re.sub(r"[^a-z0-9_]", "_", key.lower()).strip("_")
            own_id = normalized in {"id", "record_id", "entity_id", "relationship_id", "observation_id", "outcome_id", "claim_id", "formula_id", "test_id", "verse_id"}
            standard_code = "loinc" in normalized or normalized == "code" or normalized.endswith(("_code", "_key")) or normalized in {"key", "identifier", "fdc_id"}
            foreign_id = normalized.startswith(("source_", "evidence_", "parent_", "related_", "process_source_"))
            if (own_id or standard_code) and not foreign_id and not isinstance(value, dict) and str(value).strip():
                return str(value).strip()
        # Enterprise sheets conventionally put their local stable dimension in
        # column one (Scale, Hazard, Measure, System, etc.).
        for value in payload.values():
            if not isinstance(value, dict) and str(value).strip():
                return str(value).strip()
        # No invented identity: a deterministic row payload key is safer than Excel row number alone.
        return "ROW-" + stable_hash(canonical_json(payload))[:24]


def load_local(path: str | Path) -> bytes:
    return Path(path).read_bytes()
